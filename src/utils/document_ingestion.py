"""
Document Ingestion Module
Handles loading and processing of PDF, DOCX, and PPTX files
"""

import os
from typing import List
from pypdf import PdfReader
from docx import Document
from pptx import Presentation


class DocumentIngester:
    """Handles ingestion of various document formats"""
    
    def __init__(self):
        self.supported_formats = ['.pdf', '.docx', '.pptx', '.txt']
    
    def ingest_pdf(self, file_path: str) -> str:
        """Extract text from PDF file"""
        try:
            reader = PdfReader(file_path)
            text = ""
            for page in reader.pages:
                text += page.extract_text() + "\n"
            return text
        except Exception as e:
            raise Exception(f"Error reading PDF file {file_path}: {str(e)}")
    
    def ingest_docx(self, file_path: str) -> str:
        """Extract text from DOCX file"""
        try:
            doc = Document(file_path)
            text = ""
            for para in doc.paragraphs:
                text += para.text + "\n"
            return text
        except Exception as e:
            raise Exception(f"Error reading DOCX file {file_path}: {str(e)}")
    
    def ingest_pptx(self, file_path: str) -> str:
        """Extract text from PPTX file"""
        try:
            prs = Presentation(file_path)
            text = ""
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text += shape.text + "\n"
            return text
        except Exception as e:
            raise Exception(f"Error reading PPTX file {file_path}: {str(e)}")
    
    def ingest_txt(self, file_path: str) -> str:
        """Extract text from TXT file"""
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                return f.read()
        except Exception as e:
            raise Exception(f"Error reading TXT file {file_path}: {str(e)}")
    
    def ingest_document(self, file_path: str) -> dict:
        """
        Ingest a document and return its content with metadata
        
        Args:
            file_path: Path to the document file
            
        Returns:
            Dictionary with 'content' and 'metadata' keys
        """
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found: {file_path}")
        
        file_ext = os.path.splitext(file_path)[1].lower()
        
        if file_ext not in self.supported_formats:
            raise ValueError(f"Unsupported file format: {file_ext}")
        
        # Extract text based on file type
        if file_ext == '.pdf':
            content = self.ingest_pdf(file_path)
        elif file_ext == '.docx':
            content = self.ingest_docx(file_path)
        elif file_ext == '.pptx':
            content = self.ingest_pptx(file_path)
        elif file_ext == '.txt':
            content = self.ingest_txt(file_path)
        else:
            raise ValueError(f"Unsupported file format: {file_ext}")
        
        # Create metadata
        metadata = {
            'source': os.path.basename(file_path),
            'file_path': file_path,
            'file_type': file_ext,
            'file_size': os.path.getsize(file_path)
        }
        
        return {
            'content': content,
            'metadata': metadata
        }
    
    def ingest_directory(self, directory_path: str) -> List[dict]:
        """
        Ingest all supported documents from a directory
        
        Args:
            directory_path: Path to directory containing documents
            
        Returns:
            List of dictionaries with document content and metadata
        """
        if not os.path.exists(directory_path):
            raise FileNotFoundError(f"Directory not found: {directory_path}")
        
        documents = []
        for filename in os.listdir(directory_path):
            file_path = os.path.join(directory_path, filename)
            if os.path.isfile(file_path):
                file_ext = os.path.splitext(filename)[1].lower()
                if file_ext in self.supported_formats:
                    try:
                        doc = self.ingest_document(file_path)
                        documents.append(doc)
                    except Exception as e:
                        print(f"Warning: Could not ingest {filename}: {str(e)}")
        
        return documents
